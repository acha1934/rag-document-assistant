"""
Handles loading and chunking of uploaded documents.
Supports PDF, DOCX, and plain text files.
"""
import io
import uuid
from pathlib import Path
from typing import List

from langchain.schema import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter

from app.config import settings


def load_pdf(file_bytes: bytes, filename: str) -> List[Document]:
    """Extract text from a PDF file."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    docs = []
    for page_num, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": filename, "page": page_num + 1},
                )
            )
    return docs


def load_docx(file_bytes: bytes, filename: str) -> List[Document]:
    """Extract text from a DOCX file."""
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(file_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(page_content=full_text, metadata={"source": filename, "page": 1})]


def load_txt(file_bytes: bytes, filename: str) -> List[Document]:
    """Load a plain text file."""
    text = file_bytes.decode("utf-8", errors="ignore")
    return [Document(page_content=text, metadata={"source": filename, "page": 1})]


LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".doc": load_docx,
    ".txt": load_txt,
    ".md": load_txt,
}


def process_document(file_bytes: bytes, filename: str) -> tuple[str, List[Document]]:
    """
    Load, chunk, and tag a document with a unique doc_id.
    Returns (doc_id, list_of_chunks).
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in LOADERS:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: {', '.join(LOADERS.keys())}"
        )

    loader = LOADERS[suffix]
    raw_docs = loader(file_bytes, filename)

    if not raw_docs:
        raise ValueError("No text could be extracted from the document.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(raw_docs)

    doc_id = str(uuid.uuid4())
    for chunk in chunks:
        chunk.metadata["doc_id"] = doc_id
        chunk.metadata["filename"] = filename

    return doc_id, chunks
